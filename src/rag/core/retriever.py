
"""检索编排：文档入库 + 问答检索的完整流水线。

对外的核心对象。上层只管：给它一批文档 → 它入库；给它一个问题 → 它回答。
多租户：每个客户一个 Retriever 实例，doc_id 隔离。
"""
from pathlib import Path
from .config import Config
from .chunker import TextChunker, Chunk
from .embedder import Embedder
from .vector_store import VectorStore
from .generator import Generator

# 支持解析的文档后缀
_TEXT_EXTS = {".txt", ".md", ".csv", ".html", ".xml"}


def _read_text(path: Path) -> str | None:
    """按后缀路由到对应解析器。"""
    try:
        if path.suffix.lower() in _TEXT_EXTS:
            return path.read_text(encoding="utf-8", errors="ignore")
        if path.suffix.lower() == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(str(path))
            return "\n".join(p.extract_text() or "" for p in reader.pages)
        if path.suffix.lower() in {".docx"}:
            from docx import Document
            doc = Document(str(path))
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        if path.suffix.lower() == ".xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(str(path), read_only=True, data_only=True)
            rows = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c) for c in row if c is not None]
                    if cells:
                        rows.append(" ".join(cells))
            return "\n".join(rows)
    except Exception as e:
        print(f"[WARN] 解析失败 {path.name}: {e}")
    return None


class Retriever:
    def __init__(self, doc_id: str, cfg: Config | None = None):
        self.cfg = cfg or Config()
        self.doc_id = doc_id
        self.chunker = TextChunker(self.cfg.CHUNK_SIZE, self.cfg.CHUNK_OVERLAP)
        self.embedder = Embedder(self.cfg.EMBEDDING_MODEL, self.cfg.EMBEDDING_MODE)
        self.store = VectorStore(self.embedder, self.cfg.VECTOR_DB_DIR, self.doc_id)
        self.generator = Generator(self.cfg)

    # ---- 入库 ----
    def ingest_file(self, path: str | Path) -> int:
        path = Path(path)
        text = _read_text(path)
        if not text:
            return 0
        chunks = self.chunker.split(text, self.doc_id, path.name)
        self.store.add_chunks(chunks)
        self.store.save()
        return len(chunks)

    def ingest_dir(self, dirpath: str | Path, recursive: bool = True) -> int:
        dirpath = Path(dirpath)
        total = 0
        for p in (dirpath.rglob("*") if recursive else dirpath.iterdir()):
            if p.is_file():
                total += self.ingest_file(p)
        return total

    # ---- 检索 + 回答 ----
    def search(self, query: str, top_k: int | None = None) -> list:
        return self.store.search(query, top_k or self.cfg.TOP_K)

    def ask(self, question: str) -> str:
        hits = self.store.search(question, self.cfg.TOP_K)
        if not hits:
            return "（知识库暂无相关内容，请先上传文档后再提问。）"
        context = self.generator._build_context(hits)
        return self.generator.answer(question, context)

    # ---- 管理 ----
    def save(self):
        self.store.save()

    def clear(self):
        self.store.clear()
        self.save()

    @property
    def chunk_count(self) -> int:
        return self.store.count
